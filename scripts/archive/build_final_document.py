#!/usr/bin/env python3
"""Build the polished RelaySpec final research plan DOCX from its Markdown source."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

BLUE = "2E74B5"
DARK_BLUE = "17365D"
MID_BLUE = "DCE6F1"
LIGHT_BLUE = "EEF4FA"
LIGHT_GRAY = "F4F6F9"
MID_GRAY = "D9E2F3"
TEXT = RGBColor(35, 42, 49)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr, fld_char_2])


def add_hyperlink(paragraph, text: str, url: str, color=BLUE, underline=True) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(
    r"(https?://[^\s)]+|\*\*[^*]+\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))"
)


def add_inline(paragraph, text: str, size: float | None = None) -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            if size:
                run.font.size = Pt(size)
        token = match.group(0)
        if token.startswith("http"):
            add_hyperlink(paragraph, token, token)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            if size:
                run.font.size = Pt(size)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(31, 78, 121)
            run.font.size = Pt(size or 9.5)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            if size:
                run.font.size = Pt(size)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        if size:
            run.font.size = Pt(size)


def add_text_paragraph(doc, text: str, style=None, align=None, keep=False):
    p = doc.add_paragraph(style=style)
    add_inline(p, text)
    if align is not None:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_together = True
    return p


def add_callout(doc, title: str, body: str, fill=LIGHT_BLUE, accent=BLUE) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(0.10)
    table.columns[1].width = Inches(6.30)
    left, right = table.rows[0].cells
    set_cell_shading(left, accent)
    set_cell_shading(right, fill)
    set_cell_margins(left, 20, 20, 20, 20)
    set_cell_margins(right, 120, 180, 120, 180)
    p = right.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    r.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(3)
    p2 = right.add_paragraph()
    add_inline(p2, body, size=10)
    p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_equation(doc, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8FAFC")
    set_cell_margins(cell, 100, 180, 100, 180)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    eq = latex_to_readable(" ".join(line.strip() for line in lines))
    run = p.add_run(eq)
    run.font.name = "Cambria Math"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)


def latex_to_readable(text: str) -> str:
    """Convert the small LaTeX subset used by the proposal to readable math text."""
    value = text
    replacements = (
        ("\\rightarrow", "→"), ("\\mapsto", "↦"),
        ("\\left", ""), ("\\right", ""), ("\\!", ""), ("\\,", " "),
        ("\\qquad", "    "), ("\\quad", "  "),
        ("\\cap", "∩"), ("\\bigcup", "∪"),
        ("\\widetilde\\tau", "τ̃"), ("\\widehat c", "ĉ"),
        ("\\widehat q", "q̂"), ("\\mathcal L", "ℒ"),
        ("\\sum", "Σ"), ("\\prod", "Π"), ("\\Pr", "Pr"),
        ("\\lVert", "‖"), ("\\rVert", "‖"), ("\\cos", "cos"),
        ("\\log", "log"), ("\\exp", "exp"), ("\\min", "min"),
        ("\\lambda", "λ"), ("\\tau", "τ"), ("\\phi", "φ"),
        ("\\alpha", "α"), ("\\epsilon", "ε"), ("\\sigma", "σ"),
        ("\\ell", "ℓ"), ("\\cdots", "…"), ("\\ldots", "…"),
    )
    for old, new in replacements:
        value = value.replace(old, new)
    value = re.sub(r"\\operatorname\{([^{}]+)\}", r"\1", value)
    value = re.sub(r"\\mathrm\{([^{}]+)\}", r"\1", value)
    value = re.sub(r"\\mathcal\s*([A-Za-z])", r"\1", value)
    value = re.sub(r"\\widehat\s*([A-Za-z])", r"hat(\1)", value)
    value = re.sub(r"\\widetilde\s*([A-Za-z])", r"tilde(\1)", value)
    value = re.sub(r"\\boxed\{([^{}]+)\}", r"[ \1 ]", value)
    # Resolve simple fractions repeatedly; nested terms remain parenthesized text.
    simple_fraction = re.compile(r"\\frac\{([^{}]+)\}\{([^{}]+)\}")
    previous = None
    while previous != value:
        previous = value
        value = simple_fraction.sub(r"(\1)/(\2)", value)
    value = value.replace("\\frac", "/")
    value = value.replace("\\", "")
    value = value.replace("{", "(").replace("}", ")")
    value = re.sub(r"\s+", " ", value).strip()
    return value


def add_code_block(doc, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F2F2")
    set_cell_margins(cell, 110, 160, 110, 160)
    p = cell.paragraphs[0]
    p.paragraph_format.keep_together = True
    for idx, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        if idx != len(lines) - 1:
            run.add_break()


def add_markdown_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        set_row_cant_split(row)
        if row_idx == 0:
            set_repeat_table_header(row)
        for col_idx in range(cols):
            cell = row.cells[col_idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx == 0:
                set_cell_shading(cell, MID_BLUE)
            text = values[col_idx].strip() if col_idx < len(values) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, text, size=7.7 if cols >= 5 else 8.4)
            for run in p.runs:
                if row_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def choose_font(size: int, bold=False):
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def wrap(draw, text: str, font, width: int) -> list[str]:
    words = text.split()
    lines, current = [], ""
    for word in words:
        trial = f"{current} {word}".strip()
        if draw.textbbox((0, 0), trial, font=font)[2] <= width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def build_diagram(path: Path) -> None:
    image = Image.new("RGB", (1800, 880), "white")
    draw = ImageDraw.Draw(image)
    title = choose_font(42, bold=True)
    label = choose_font(28, bold=True)
    box_label = choose_font(22, bold=True)
    body = choose_font(20)
    small = choose_font(20)

    draw.text((70, 42), "RelaySpec removes repeated source-model execution", font=title, fill="#17365D")
    draw.text((70, 118), "Naive cross-target reuse", font=label, fill="#7F1D1D")
    draw.text((70, 488), "RelaySpec", font=label, fill="#166534")

    def box(x, y, w, h, fill, outline, heading, content):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=22, fill=fill, outline=outline, width=4)
        heading_lines = wrap(draw, heading, box_label, w - 48)
        heading_y = y + 20
        for heading_line in heading_lines:
            draw.text((x + 24, heading_y), heading_line, font=box_label, fill=outline)
            heading_y += 29
        yy = max(y + 70, heading_y + 8)
        for line in wrap(draw, content, body, w - 48):
            draw.text((x + 24, yy), line, font=body, fill="#263238")
            yy += 28

    def arrow(x1, y, x2, color="#5B6573"):
        draw.line((x1, y, x2, y), fill=color, width=7)
        draw.polygon([(x2, y), (x2 - 20, y - 14), (x2 - 20, y + 14)], fill=color)

    box(70, 175, 350, 205, "#FEE2E2", "#991B1B", "Old target A trunk", "Runs every cycle to recreate the proposer interface (20.45 ms in the current probe).")
    box(530, 175, 310, 205, "#FFF7ED", "#9A3412", "Frozen proposer A", "Generates a block using the interface learned from target A.")
    box(950, 175, 330, 205, "#EFF6FF", "#1D4ED8", "New target B", "Verifies the proposed block exactly and computes B hidden states.")
    box(1390, 175, 330, 205, "#F3F4F6", "#374151", "Commit", "Accept the verified prefix and correction token; crop rejected cache entries.")
    arrow(420, 278, 530); arrow(840, 278, 950); arrow(1280, 278, 1390)

    box(70, 545, 350, 205, "#EFF6FF", "#1D4ED8", "Target B hidden taps", "Reuse representations already produced during exact verification.")
    box(530, 545, 310, 205, "#DCFCE7", "#166534", "Relay B to A", "Translate B features into proposer A's interface (0.304 ms in the current probe).")
    box(950, 545, 330, 205, "#FFF7ED", "#9A3412", "Same frozen proposer A", "No proposer architecture or weight changes; propose and verify normally.")
    box(1390, 545, 330, 205, "#EFF6FF", "#1D4ED8", "Exact target B commit", "Verifier B remains authoritative; proposal errors affect speed, not greedy output.")
    arrow(420, 648, 530, "#166534"); arrow(840, 648, 950, "#166534"); arrow(1280, 648, 1390, "#166534")

    draw.rounded_rectangle((635, 790, 1165, 842), radius=18, fill="#D1FAE5", outline="#166534", width=3)
    draw.text((680, 802), "Old target A transformer trunk removed", font=small, fill="#166534")
    image.save(path, quality=95)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.22

    for name, size, color, before, after in (
        ("Title", 28, DARK_BLUE, 0, 16),
        ("Heading 1", 17, BLUE, 18, 9),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, DARK_BLUE, 9, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.3)
        style.paragraph_format.left_indent = Inches(0.36)
        style.paragraph_format.first_line_indent = Inches(-0.19)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.15


def add_header_footer(doc: Document) -> None:
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.text = "RELAYSPEC  |  FINAL RESEARCH PROPOSAL AND EXECUTION PLAN"
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(BLUE)
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("RelaySpec  •  26 August 2026  •  ")
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(90, 100, 110)
        add_page_number(p)


def add_cover(doc: Document) -> None:
    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = band.cell(0, 0)
    set_cell_shading(cell, DARK_BLUE)
    set_cell_margins(cell, 80, 100, 80, 100)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FINAL RESEARCH PROPOSAL  •  ICLR EVIDENCE PLAN")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(255, 255, 255)

    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("RelaySpec")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(34)
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Frozen-Interface Emulation for Portable Speculative Proposers")
    r.font.name = "Calibri"
    r.font.size = Pt(17)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p.paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Method • Theory • Novelty audit • Experiments • Implementation • Paper blueprint")
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(80, 90, 100)

    doc.add_paragraph()
    add_callout(
        doc,
        "Central question",
        "Can a small translator reuse a frozen proposer trained for target A on a compatible target B, remove A's repeated transformer trunk, recover most native-proposer acceptance, and accelerate exact target-B decoding?",
        fill="EAF2F8",
        accent=DARK_BLUE,
    )

    cards = doc.add_table(rows=1, cols=3)
    cards.alignment = WD_TABLE_ALIGNMENT.CENTER
    cards.autofit = False
    values = [
        ("1.2408x", "measured over naive cross-target reuse"),
        ("12 / 12", "greedy sequence matches in the current probe"),
        ("Not measured", "native Qwen3-14B AR comparison"),
    ]
    for idx, (value, caption) in enumerate(values):
        cell = cards.cell(0, idx)
        set_cell_shading(cell, LIGHT_GRAY if idx != 1 else LIGHT_BLUE)
        set_cell_margins(cell, 130, 110, 130, 110)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        r.bold = True
        r.font.size = Pt(16 if idx != 2 else 13)
        r.font.color.rgb = RGBColor.from_string(BLUE if idx != 2 else "9A3412")
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(caption)
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = RGBColor(80, 90, 100)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("26 August 2026  |  Research-ready revision")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(90, 100, 110)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_document_map(doc: Document, headings: list[str]) -> None:
    doc.add_heading("Document map", level=1)
    add_text_paragraph(
        doc,
        "This document is the canonical RelaySpec proposal and execution plan. It incorporates the external novelty feedback, separates measured evidence from targets, and orders experiments by decisive value per GPU-hour.",
    )
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, heading in enumerate(headings, start=1):
        row = table.add_row()
        set_row_cant_split(row)
        set_cell_shading(row.cells[0], MID_BLUE)
        set_cell_margins(row.cells[0])
        set_cell_margins(row.cells[1])
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(f"{idx:02d}")
        r0.bold = True
        r0.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        add_inline(row.cells[1].paragraphs[0], heading, size=9)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def clean_table_cell(value: str) -> str:
    return value.strip().strip("|").strip()


def parse_markdown(doc: Document, source: str, diagram_path: Path) -> None:
    lines = source.splitlines()
    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped or stripped == "---":
            index += 1
            continue
        if stripped.startswith(
            ("# RelaySpec:", "**Final research", "**Date:", "**Status:")
        ):
            index += 1
            continue
        if stripped == "<!-- SYSTEM_DIAGRAM -->":
            picture = doc.add_picture(str(diagram_path), width=Inches(6.55))
            picture._inline.docPr.set(
                "descr",
                "RelaySpec replaces repeated execution of source target A with a small relay from target B hidden states to the frozen proposer A interface; exact target B verification remains authoritative.",
            )
            picture._inline.docPr.set("title", "RelaySpec system diagram")
            p = doc.add_paragraph("Figure 1. The source trunk removed by RelaySpec; target verification remains authoritative.")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.italic = True
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(90, 100, 110)
            index += 1
            continue
        if stripped.startswith("```"):
            code = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index].rstrip())
                index += 1
            add_code_block(doc, code)
            index += 1
            continue
        if stripped == "\\[":
            equation = []
            index += 1
            while index < len(lines) and lines[index].strip() != "\\]":
                equation.append(lines[index])
                index += 1
            add_equation(doc, equation)
            index += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
            index += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
            index += 1
            continue
        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:], level=3)
            index += 1
            continue
        if stripped.startswith("| "):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = []
            for table_line in table_lines:
                values = [clean_table_cell(v) for v in table_line.strip("|").split("|")]
                if values and all(re.fullmatch(r":?-{3,}:?", value) for value in values):
                    continue
                rows.append(values)
            add_markdown_table(doc, rows)
            continue
        if re.match(r"^- ", stripped):
            content = stripped[2:]
            index += 1
            while index < len(lines):
                continuation = lines[index]
                nxt = continuation.strip()
                if not nxt or not continuation[:1].isspace() or re.match(r"^(?:-|\d+\.)\s", nxt):
                    break
                content += " " + nxt
                index += 1
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, content)
            continue
        ordered = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if ordered:
            number, content = ordered.group(1), ordered.group(2)
            index += 1
            while index < len(lines):
                continuation = lines[index]
                nxt = continuation.strip()
                if not nxt or not continuation[:1].isspace() or re.match(r"^(?:-|\d+\.)\s", nxt):
                    break
                content += " " + nxt
                index += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.38)
            p.paragraph_format.first_line_indent = Inches(-0.28)
            p.paragraph_format.space_after = Pt(3)
            prefix = p.add_run(f"{number}.  ")
            prefix.font.size = Pt(10.3)
            add_inline(p, content)
            continue
        if stripped.startswith("> "):
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().lstrip(">").strip())
                index += 1
            add_callout(doc, "Core thesis", " ".join(quote_lines), fill="EAF2F8", accent=DARK_BLUE)
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            nxt = lines[index].strip()
            if (
                not nxt
                or nxt == "---"
                or nxt.startswith(("#", "|", "- ", ">", "```", "<!--"))
                or nxt == "\\["
                or re.match(r"^\d+\. ", nxt)
            ):
                break
            paragraph_lines.append(nxt)
            index += 1
        text = " ".join(paragraph_lines)
        add_text_paragraph(doc, text)


def build(source_path: Path, output_path: Path) -> None:
    source = source_path.read_text(encoding="utf-8")
    headings = [line[3:].strip() for line in source.splitlines() if line.startswith("## ")]
    assets_dir = output_path.parent / ".relayspec_document_assets"
    assets_dir.mkdir(parents=True, exist_ok=True)
    diagram_path = assets_dir / "relayspec-system-diagram.png"
    build_diagram(diagram_path)

    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_document_map(doc, headings)
    parse_markdown(doc, source, diagram_path)
    add_header_footer(doc)

    core = doc.core_properties
    core.title = "RelaySpec: Frozen-Interface Emulation for Portable Speculative Proposers"
    core.subject = "Final research proposal, implementation plan, and ICLR evidence blueprint"
    core.author = "RelaySpec research project"
    core.keywords = "speculative decoding, DFlash, EAGLE-3, inference efficiency, interface emulation"
    core.comments = "Generated from the canonical Markdown research plan."

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> int:
    if len(sys.argv) != 3:
        print("usage: build_final_document.py SOURCE.md OUTPUT.docx", file=sys.stderr)
        return 2
    build(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
